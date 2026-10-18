import os
import sys
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, BASE_DIR)

from src.analytics.kpi_calculator import get_fleet_kpis, get_budget_variance
from src.database.db_manager import query_dataframe


def generate_fleet_excel_report(output_path=None):
    """
    Menghasilkan workbook Microsoft Excel korporat dengan 
    formula native Excel (=SUM, =AVERAGE, =IF) dan visual kartu KPI yang rapi
    """
    wb = openpyxl.Workbook()
    
    # Definisi Gaya, Font & Warna
    font_title = Font(name="Calibri", size=15, bold=True, color="1F4E79")
    font_subtitle = Font(name="Calibri", size=10, italic=True, color="595959")
    font_section = Font(name="Calibri", size=11, bold=True, color="1F4E79")
    
    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=10)
    font_bold = Font(name="Calibri", size=10, bold=True)
    
    font_card_title = Font(name="Calibri", size=9, bold=True, color="595959")
    font_card_num = Font(name="Calibri", size=16, bold=True, color="1F4E79")
    font_card_alert = Font(name="Calibri", size=16, bold=True, color="C00000")

    fill_navy = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    fill_card = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_card_alert = PatternFill(start_color="FFF1F2", end_color="FFF1F2", fill_type="solid")
    fill_success = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
    fill_alert = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")

    thin_side = Side(style="thin", color="D1D5DB")
    thick_bottom = Side(style="double", color="1F4E79")
    
    border_all = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    border_total = Border(top=thin_side, bottom=thick_bottom, left=thin_side, right=thin_side)

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # Sheet 1: Executive Summary
    ws_kpi = wb.active
    ws_kpi.title = "Executive Summary"
    ws_kpi.views.sheetView[0].showGridLines = True

    # Title Banner (Kolom B s/d G)
    ws_kpi["B2"] = "HEAVY TRUCK FLEET OPERATIONAL PERFORMANCE REPORT"
    ws_kpi["B2"].font = font_title
    ws_kpi["B3"] = "Executive Intelligence: Fleet Availability, Reliability MTBF & Maintenance Budget Variance"
    ws_kpi["B3"].font = font_subtitle

    kpis = get_fleet_kpis(days=30)

    # Helper untuk menggambar kotak kartu KPI yang simetris
    def draw_card(ws, start_col, end_col, label, value, is_alert=False):
        fill_bg = fill_card_alert if is_alert else fill_card
        num_font = font_card_alert if is_alert else font_card_num
        
        # Merge header (Row 5)
        ws.merge_cells(f"{start_col}5:{end_col}5")
        top_cell = ws[f"{start_col}5"]
        top_cell.value = label
        top_cell.font = font_card_title
        top_cell.alignment = align_center
        top_cell.fill = fill_bg

        # Merge value (Row 6)
        ws.merge_cells(f"{start_col}6:{end_col}6")
        val_cell = ws[f"{start_col}6"]
        val_cell.value = value
        val_cell.font = num_font
        val_cell.alignment = align_center
        val_cell.fill = fill_bg

        # Beri border keliling
        cols = [start_col] if start_col == end_col else [start_col, end_col]
        for col in cols:
            for r in [5, 6]:
                ws[f"{col}{r}"].border = border_all
                ws[f"{col}{r}"].fill = fill_bg

    # 4 Kotak KPI Cards yang lurus dengan tabel di bawahnya (Cols B s/d G)
    draw_card(ws_kpi, "B", "C", "FLEET AVAILABILITY (PA)", f"{kpis['physical_availability_pct']:.2f}%")
    draw_card(ws_kpi, "D", "D", "MTBF (HOURS)", f"{kpis['mtbf_hours']:,.1f} hrs")
    draw_card(ws_kpi, "E", "E", "BREAKDOWNS", f"{kpis['total_breakdowns']} units", is_alert=True)
    draw_card(ws_kpi, "F", "G", "TOTAL MAINT. COST", f"${kpis['total_maintenance_cost_usd']:,.2f}")

    ws_kpi.row_dimensions[5].height = 20
    ws_kpi.row_dimensions[6].height = 28

    # Section 1: Tabel Perencanaan Anggaran Bulanan
    ws_kpi["B8"] = "1. ANALISIS ANGGARAN PERAWATAN BULANAN (BUDGET VARIANCE)"
    ws_kpi["B8"].font = font_section

    headers_budget = ["Periode", "Anggaran (USD)", "Aktual Biaya (USD)", "Varians Budget (USD)", "Target PA", "Status Anggaran"]
    cols_budget = ["B", "C", "D", "E", "F", "G"]

    for col, h in zip(cols_budget, headers_budget):
        cell = ws_kpi[f"{col}9"]
        cell.value = h
        cell.font = font_header
        cell.fill = fill_navy
        cell.alignment = align_center
        cell.border = border_all
    ws_kpi.row_dimensions[9].height = 24

    budget_df = get_budget_variance()
    start_row = 10

    for idx, row in budget_df.iterrows():
        r = start_row + idx
        ws_kpi[f"B{r}"] = str(row["period_month"])
        ws_kpi[f"C{r}"] = float(row["allocated_budget_usd"])
        ws_kpi[f"D{r}"] = float(row["actual_cost_usd"])
        
        # Formula native: Varians = Budget - Actual
        ws_kpi[f"E{r}"] = f"=C{r}-D{r}"
        
        # Target PA sebagai nilai persentase numerik (0.92, 0.93, 0.94)
        ws_kpi[f"F{r}"] = float(row["target_fleet_availability"]) / 100.0
        
        # Formula native: Status
        ws_kpi[f"G{r}"] = f'=IF(E{r}>=0, "WITHIN BUDGET", "OVER BUDGET")'

        ws_kpi[f"B{r}"].alignment = align_center
        ws_kpi[f"C{r}"].alignment = align_right
        ws_kpi[f"D{r}"].alignment = align_right
        ws_kpi[f"E{r}"].alignment = align_right
        ws_kpi[f"F{r}"].alignment = align_center
        ws_kpi[f"G{r}"].alignment = align_center

        ws_kpi[f"C{r}"].number_format = "$#,##0.00"
        ws_kpi[f"D{r}"].number_format = "$#,##0.00"
        ws_kpi[f"E{r}"].number_format = "$#,##0.00"
        ws_kpi[f"F{r}"].number_format = "0.0%"

        # Conditional highlight untuk status
        ws_kpi[f"G{r}"].fill = fill_success if (row["allocated_budget_usd"] >= row["actual_cost_usd"]) else fill_alert

        for c in cols_budget:
            ws_kpi[f"{c}{r}"].border = border_all
            ws_kpi[f"{c}{r}"].font = font_data
        ws_kpi.row_dimensions[r].height = 20

    # Baris TOTAL
    tot_row = start_row + len(budget_df)
    ws_kpi[f"B{tot_row}"] = "TOTAL"
    ws_kpi[f"C{tot_row}"] = f"=SUM(C{start_row}:C{tot_row-1})"
    ws_kpi[f"D{tot_row}"] = f"=SUM(D{start_row}:D{tot_row-1})"
    ws_kpi[f"E{tot_row}"] = f"=SUM(E{start_row}:E{tot_row-1})"
    
    # Rata-rata Target PA yang benar (=AVERAGE kolom F!)
    ws_kpi[f"F{tot_row}"] = f"=AVERAGE(F{start_row}:F{tot_row-1})"
    ws_kpi[f"G{tot_row}"] = ""

    ws_kpi[f"B{tot_row}"].alignment = align_center
    ws_kpi[f"C{tot_row}"].alignment = align_right
    ws_kpi[f"D{tot_row}"].alignment = align_right
    ws_kpi[f"E{tot_row}"].alignment = align_right
    ws_kpi[f"F{tot_row}"].alignment = align_center

    ws_kpi[f"C{tot_row}"].number_format = "$#,##0.00"
    ws_kpi[f"D{tot_row}"].number_format = "$#,##0.00"
    ws_kpi[f"E{tot_row}"].number_format = "$#,##0.00"
    ws_kpi[f"F{tot_row}"].number_format = "0.0%"

    for c in cols_budget:
        ws_kpi[f"{c}{tot_row}"].font = font_bold
        ws_kpi[f"{c}{tot_row}"].border = border_total
    ws_kpi.row_dimensions[tot_row].height = 22

    # Section 2: Ringkasan Biaya Mogok vs Servis Preventif ($10 vs $500)
    sec2_title_row = tot_row + 2
    ws_kpi[f"B{sec2_title_row}"] = "2. STRUKTUR BIAYA: SERVIS PREVENTIF ($10) VS MOGOK DARURAT ($500)"
    ws_kpi[f"B{sec2_title_row}"].font = font_section

    sec2_header_row = sec2_title_row + 1
    headers_sec2 = ["Kategori Servis", "Unit Biaya (Tarif)", "Jumlah Insiden", "Total Biaya (USD)", "Porsi Biaya (%)", "Status Dampak"]
    for col, h in zip(cols_budget, headers_sec2):
        cell = ws_kpi[f"{col}{sec2_header_row}"]
        cell.value = h
        cell.font = font_header
        cell.fill = fill_navy
        cell.alignment = align_center
        cell.border = border_all
    ws_kpi.row_dimensions[sec2_header_row].height = 24

    # Data Mogok Darurat vs Preventif
    r_brk = sec2_header_row + 1
    r_prv = sec2_header_row + 2
    r_tot_sec2 = sec2_header_row + 3

    # Row Breakdown
    ws_kpi[f"B{r_brk}"] = "Mogok Darurat di Jalan (Roadside Failure)"
    ws_kpi[f"C{r_brk}"] = "$500.00"
    ws_kpi[f"D{r_brk}"] = int(kpis["total_breakdowns"])
    ws_kpi[f"E{r_brk}"] = f"=D{r_brk}*500"
    ws_kpi[f"F{r_brk}"] = f"=E{r_brk}/E{r_tot_sec2}"
    ws_kpi[f"G{r_brk}"] = "BIAYA TINGGI (Target ML)"

    # Row Preventive
    ws_kpi[f"B{r_prv}"] = "Inspeksi Preventif Terencana di Bengkel"
    ws_kpi[f"C{r_prv}"] = "$10.00"
    ws_kpi[f"D{r_prv}"] = int(kpis["total_preventive_inspections"])
    ws_kpi[f"E{r_prv}"] = f"=D{r_prv}*10"
    ws_kpi[f"F{r_prv}"] = f"=E{r_prv}/E{r_tot_sec2}"
    ws_kpi[f"G{r_prv}"] = "EFISIEN (Hemat)"

    # Total Row
    ws_kpi[f"B{r_tot_sec2}"] = "TOTAL BIAYA PERAWATAN"
    ws_kpi[f"C{r_tot_sec2}"] = "-"
    ws_kpi[f"D{r_tot_sec2}"] = f"=SUM(D{r_brk}:D{r_prv})"
    ws_kpi[f"E{r_tot_sec2}"] = f"=SUM(E{r_brk}:E{r_prv})"
    ws_kpi[f"F{r_tot_sec2}"] = f"=SUM(F{r_brk}:F{r_prv})"
    ws_kpi[f"G{r_tot_sec2}"] = ""

    # Formatting Section 2
    for r, is_tot, is_alert in [(r_brk, False, True), (r_prv, False, False), (r_tot_sec2, True, False)]:
        ws_kpi[f"B{r}"].alignment = align_left
        ws_kpi[f"C{r}"].alignment = align_center
        ws_kpi[f"D{r}"].alignment = align_center
        ws_kpi[f"E{r}"].alignment = align_right
        ws_kpi[f"F{r}"].alignment = align_center
        ws_kpi[f"G{r}"].alignment = align_center

        ws_kpi[f"E{r}"].number_format = "$#,##0.00"
        ws_kpi[f"F{r}"].number_format = "0.0%"

        border_use = border_total if is_tot else border_all
        font_use = font_bold if is_tot else font_data
        
        for c in cols_budget:
            ws_kpi[f"{c}{r}"].border = border_use
            ws_kpi[f"{c}{r}"].font = font_use
        
        if not is_tot:
            ws_kpi[f"G{r}"].fill = fill_alert if is_alert else fill_success
        ws_kpi.row_dimensions[r].height = 20

    # Lebar kolom spesifik yang pas dan proporsional untuk Sheet 1
    col_widths = {
        "A": 4,
        "B": 38,
        "C": 20,
        "D": 22,
        "E": 22,
        "F": 18,
        "G": 24
    }
    for col_l, w in col_widths.items():
        ws_kpi.column_dimensions[col_l].width = w

    # -------------------------------------------------------------------------
    # Sheet 2: Log Work Orders Servis
    # -------------------------------------------------------------------------
    ws_maint = wb.create_sheet(title="Maintenance Work Orders")
    ws_maint.views.sheetView[0].showGridLines = True

    ws_maint["A1"] = "LOG WORK ORDER PERBAIKAN & INSPEKSI ARMADA"
    ws_maint["A1"].font = font_title
    ws_maint["A2"] = "Tracking Servis Darurat ($500) vs Inspeksi Preventif Rutin ($10)"
    ws_maint["A2"].font = font_subtitle

    headers = ["Record ID", "Truck ID", "Tanggal", "Tipe Servis", "Subsistem", "Downtime (Jam)", "Biaya Riil (USD)", "Catatan Mekanik"]
    for col_idx, h in enumerate(headers, 1):
        cell = ws_maint.cell(row=4, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_navy
        cell.alignment = align_center
        cell.border = border_all
    ws_maint.row_dimensions[4].height = 24

    maint_df = query_dataframe("SELECT * FROM maintenance_records ORDER BY inspection_date DESC LIMIT 50;")
    m_start = 5

    for idx, row in maint_df.iterrows():
        curr_row = m_start + idx
        ws_maint.cell(row=curr_row, column=1, value=int(row["record_id"])).alignment = align_center
        ws_maint.cell(row=curr_row, column=2, value=str(row["truck_id"])).alignment = align_center
        ws_maint.cell(row=curr_row, column=3, value=str(row["inspection_date"])).alignment = align_center
        
        st_cell = ws_maint.cell(row=curr_row, column=4, value=str(row["service_type"]))
        st_cell.alignment = align_center
        
        if row["service_type"] == "EMERGENCY_BREAKDOWN":
            st_cell.fill = fill_alert
        else:
            st_cell.fill = fill_success

        ws_maint.cell(row=curr_row, column=5, value=str(row["failure_subsystem"])).alignment = align_center
        ws_maint.cell(row=curr_row, column=6, value=float(row["downtime_hours"])).alignment = align_right
        
        cost_cell = ws_maint.cell(row=curr_row, column=7, value=float(row["actual_cost_usd"]))
        cost_cell.number_format = "$#,##0.00"
        cost_cell.alignment = align_right

        ws_maint.cell(row=curr_row, column=8, value=str(row["mechanic_notes"])).alignment = align_left

        for c_idx in range(1, 9):
            ws_maint.cell(row=curr_row, column=c_idx).border = border_all
            ws_maint.cell(row=curr_row, column=c_idx).font = font_data
        ws_maint.row_dimensions[curr_row].height = 19

    # Lebar kolom Sheet 2
    sheet2_widths = {
        "A": 12, "B": 14, "C": 14, "D": 26, "E": 24, "F": 16, "G": 18, "H": 55
    }
    for col_l, w in sheet2_widths.items():
        ws_maint.column_dimensions[col_l].width = w

    # Simpan file fisik
    if output_path:
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        wb.save(output_path)

    byte_stream = io.BytesIO()
    wb.save(byte_stream)
    byte_stream.seek(0)
    return byte_stream.getvalue()


if __name__ == "__main__":
    report_file = os.path.join(BASE_DIR, "data", "Fleet_Performance_Report.xlsx")
    generate_fleet_excel_report(report_file)
    print("Laporan Excel korporat berhasil diperbarui dengan layout rapi dan simetris!")
    print("Lokasi file:", report_file)